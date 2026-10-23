"""用 Codex 配套 Python 生成文档；实验环境不需要安装 python-docx。"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent


def font(style,name='宋体',size=11):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)


def build(source,destination):
    document=Document();section=document.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(2.2);section.right_margin=Cm(2.2)
    normal=document.styles['Normal'];font(normal,size=11)
    normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12)]:
        style=document.styles[name];font(style,'黑体',size)
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.keep_with_next=True
    footer=section.footer.paragraphs[0];footer.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    text=source.read_text(encoding='utf-8')
    for block in text.split('\n\n'):
        block=block.strip('\n')
        if not block.strip() or block.strip()=='---PAGE---':continue
        if block.startswith('# '):document.add_paragraph(block[2:],'Title')
        elif block.startswith('## '):document.add_paragraph(block[3:],'Heading 1')
        elif block.startswith('### '):document.add_paragraph(block[4:],'Heading 2')
        elif block.startswith('    '):
            for line in block.splitlines():
                p=document.add_paragraph();p.paragraph_format.space_after=Pt(3)
                run=p.add_run(line.strip());run.font.name='Consolas';run.font.size=Pt(8.5)
        else:
            p=document.add_paragraph()
            parts=re.split(r'(\*\*.*?\*\*)',block.replace('\n',' '))
            for part in parts:
                run=p.add_run(part[2:-2] if part.startswith('**') else part.replace('`',''))
                run.bold=part.startswith('**')
    document.core_properties.title=source.stem;document.core_properties.subject='TSRD 雷达信号分选实验'
    document.core_properties.author='';document.core_properties.keywords='TSRD, PDW, benchmark'
    document.save(destination)
    print(str(destination),flush=True)


if __name__=='__main__':
    for name in ('数据生成与评估方案','代码说明'):
        build(ROOT/(name+'.md'),ROOT/(name+'.docx'))
