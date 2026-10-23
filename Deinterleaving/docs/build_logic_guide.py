"""从当前工程逻辑生成独立的中文代码导读 Word，不读取或修改原始数据。"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "雷达分选Benchmark代码逻辑导读.docx"


def set_font(style, chinese="等线", western="Aptos", size=10.5, bold=False):
    style.font.name = western
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(25, 35, 45)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), chinese)
    fonts.set(qn("w:ascii"), western)
    fonts.set(qn("w:hAnsi"), western)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin = section.right_margin = Cm(2.2)
section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)

normal = doc.styles["Normal"]
set_font(normal)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.17
for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 11.5)):
    style = doc.styles[name]
    set_font(style, chinese="等线", size=size, bold=True)
    style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
    if name == "Title":
        # Word 内建 Title 可能带蓝色底边，清掉该样式的段落边框。
        ppr = style.element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)

title = doc.add_paragraph("雷达分选 Benchmark 代码逻辑导读", "Title")
title_ppr = title._p.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)
p = doc.add_paragraph("阅读对象是要理解、修改或在新机器运行 Deinterleaving 工程的人。本文按一次实验真正执行的顺序，解释数据怎样变成样本、算法怎样产生标签、指标怎样计算，以及每类方法目前能验证到什么程度。所有描述对应当前代码，而非预期中的未来实现。")
p.paragraph_format.space_after = Pt(9)


def heading(text, level=1):
    return doc.add_paragraph(text, "Heading 1" if level == 1 else "Heading 2")


def para(text):
    doc.add_paragraph(text)


def lines(items):
    for label, body in items:
        p = doc.add_paragraph(style="Normal")
        r = p.add_run(label + "  ")
        r.bold = True
        p.add_run(body)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                    fonts.set(qn("w:eastAsia"), "等线")
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    # 表头在分页后重复，整行不跨页断开。
    for j, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        if j == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


heading("一  整体执行链路")
para("一次评估由 benchmark.py 的 evaluate 命令驱动。其执行链可读成：读取配置 → 拟合或载入对应模式的 train 统计 → 检查所选 HDF5 文件 → 确定随机窗口 → 生成场景样本 → 对同一份样本逐个运行算法 → 用样本真值评分 → 保存逐项记录 → 按文件汇总。训练是另一条链：train 原始窗口 → val 选最优轮次 → 冻结权重 → evaluate 加载权重。")
para("这两条链以 profile 和 checkpoint 相连。算法统一只接收观测矩阵 x、训练统计 profile、参数和种子；真实标签 y、父脉冲编号 parent、场景名称以及真实源数不传给 predict。真值只用于规定的场景构造、密集场景筛选与评分。")
table(["位置", "主要职责", "关键入口"], [
    ("benchmark.py", "配置、全流程编排、运行快照、断点续跑、源文件检查", "main / run_evaluation"),
    ("data", "六子集 HDF5 读取、统计、窗口和场景变换", "loader / inventory / scenarios"),
    ("algorithms", "25 个名称的注册、基线和论文组件预测", "registry.predict"),
    ("train.py", "三个深度基线的 train 与 val 训练", "train_models"),
    ("evaluation", "隔离运行、逐样本评分、文件等权汇总", "runner / metrics / report"),
    ("configs 与 results", "实验参数和每次运行的证据", "JSON 配置与输出目录"),
], [3.0, 9.7, 3.9])

heading("二  数据从 HDF5 到算法输入")
heading("六个子集与文件筛选", 2)
para("data/loader.py 的 expand_subsets 只接受 scan/train、scan/val、scan/test、stare/train、stare/val、stare/test 或 all。files_for 将逻辑名映射到实际目录，例如 scan/test 对应 scan/test_scan，并按 config 编号排序。配置可指定文件名、文件数、观察时间和窗口数；多选子集仍独立汇总。")
heading("字段核验与单窗口读取", 2)
para("feature_order 读取 HDF5 的 metadata/feature_names，确认五列是 ToA、Frequency、PulseWidth、AoA、Amplitude 后才调整顺序。load_window 用只读模式打开文件，以半开区间 [start, start+duration) 选 PDW；有序 ToA 使用连续切片，无序文件按索引筛选。时间减去窗口起点并稳定排序。内部 x 为 float64 的 N×5 矩阵，单位依次为 μs、MHz、μs、度、dB；y 为 int64 源编号，-1 表示虚警；parent 是原始脉冲行号。")
para("窗口起点由子集、文件名、种子共同确定，允许读到空窗口。短观察场景再从普通窗口起点截取 100 至 300 μs，所以不会为短窗口重新挑一个更容易的起点。window_us=null 则读取整条记录。")
heading("训练统计和特征", 2)
para("fit_train_profile 仅从对应接收模式的 train 文件拟合各特征的中位数、四分位距、0.1% 与 99.9% 分位范围，并计算正相邻 ToA 差的中位数作为时间误差尺度。这个时间尺度不被解释为真实发射 PRI。pdw_features 对 RF、PW、幅度使用训练统计标准化；AOA 转成 sin/cos，避免 -180 度和 180 度被错误视为远离。时间模型额外输入 log(1+DTOA/时间尺度)，空间聚类默认只用 RF、PW 与角度两维，不使用幅度或绝对 ToA。")
heading("密集多源与低 SNR", 2)
para("inventory 实际读取 labels，排除 -1 后统计各 HDF5 的真实源数；统计 CSV 只辅助提供时间范围，脉冲数不一致时改用 HDF5 实测值。dense_selection 在当前子集和当前文件筛选范围中找源数至少 120 的全部记录；若没有，保留源数最大的全部并列记录。先选记录，再读窗口，因此结果同时记原记录源数与窗口可见源数。data/snr.py 可读外部带 snr_db 的 NPZ；现有 TSRD 无对应标注，场景直接记未评估。")

heading("三  场景样本如何生成")
para("data/scenarios.py 的 transform 先复制 Sample，再按照场景修改内存数据，最后按 ToA 稳定排序并记录生成数量和参数。默认 grid 是无扰动基线加单因素网格，共 64 个设置；compose 允许依指定顺序叠加多个变换，此时每一步的比例分母是该步输入。相同源窗口和变换种子可供 25 个算法共用。")
table(["场景", "代码执行规则"], [
    ("缺失 10% 至 90%", "从 y≥0 的真实脉冲无放回删除；k=floor(pN+0.5)，记录消失源。"),
    ("虚警 10% 至 50%", "增加 k=floor(pN+0.5) 个 y=-1 的 PDW；ToA 在窗口均匀采样，其余列按 train 范围采样。"),
    ("参数噪声 1% 至 10%", "各列加独立高斯误差；RF/PW 用相对尺度，ToA 用 train 时间尺度，AOA/幅度用 train 四分位尺度；再裁剪时间和正值、环绕角度。"),
    ("PDW 重叠 10% 至 90%", "按训练尺度的 RF/PW/圆周角距离选最近两源，使两源中心向共同中点移动，源内差值和 ToA 保留。少于两源记不适用。"),
    ("分裂 5% 至 30%", "选中父脉冲按 0.3 至 0.7 的比例拆成连续两个片段；两片同源、同 parent，两个脉宽之和等于原脉宽。"),
    ("短观察 100 至 300 μs", "从同一普通窗口左端截取，保留空窗口与极少脉冲窗口。"),
    ("密集多源与低 SNR", "前者是记录筛选，不改 PDW；后者当前缺少真实 SNR 输入，返回未评估。"),
], [4.0, 12.6])
heading("PRI 与 RF 调制", 2)
para("PRI 调制逐源重新生成 ToA，使用 t[n+1]=t[n]+PRI[n]，默认合成基准 B=100 μs，随机初相；生成后裁掉超出原窗口的脉冲，并重排所有源。抖动为 B 的 ±10% 独立均匀变化；参差循环 [0.8,1,1.2]B；滑变在 64 个间隔内从 0.8B 线性到 1.2B 后保持，可改为下降；组变每 16 个间隔换一档；随机独立取 [0.5,1.5]B；驻留随机维持 32 至 64 个间隔后换档；交错将 B 与 1.3B 两条不同相位的子序列合并为同一源。这里的 B 是显式合成参数，不由 scan 漏检序列反推。")
para("RF 调制逐源取原 RF 中位数，默认在该中心 ±10% 与 train 合法上界之间变化。固定保持中心频率；随机固定每源开头取一次值；范围捷变逐脉冲独立取值；线性在默认 64 脉冲到达上界后保持；锯齿每 64 脉冲复位；同时多频复制同一发射事件，令副本共用 ToA、源编号与 parent，但频率不同。PRI 与 RF 重生成并不重跑原接收机扫描过程。")

heading("四  算法的统一输出")
para("algorithms/registry.py 注册 25 个名称。predict 必须对每个输入 PDW 返回一个整数标签，非负整数表示算法分出的簇，-1 表示拒识。编号本身没有物理意义，评分时通过匹配处理任意标签置换。registry 还为论文方法保存 fidelity、主榜或先验榜 track、是否完成严格复现验证。正式配置 strict_papers_only=true，会让全部尚未核验的论文项明确返回未评估。")
heading("六种 TOA 基线", 2)
para("algorithms/classical.py 先从混合 ToA 提出候选 PRI，再用 sequence_chain 动态规划寻找容许误差与漏检整数倍的最长时间链；已分配的脉冲不再分给下一条链。下表区别候选生成方式，后端序列抽取共用。")
table(["方法", "候选 PRI 逻辑"], [
    ("CDIF", "累加多个差分阶数的 ToA 差直方图并找峰。"),
    ("SDIF", "每个差分阶数先过独立阈值，再汇集保留的峰。"),
    ("PRI Transform", "按差分 PRI 分箱，累加到达时刻的复相位，再取幅值找峰。"),
    ("Hough", "先取高频差分候选，在 PRI 与相位平面上投票。"),
    ("Search", "直接从相邻 1 至 5 阶时间差提出候选，再做序列搜索。"),
    ("Correlation", "把 ToA 栅格化成时间序列，用 FFT 自相关找周期峰。"),
], [3.4, 13.2])
heading("六种空间聚类基线", 2)
para("K-means、DBSCAN、HDBSCAN、Mean Shift、Spectral、Hierarchical 都在训练尺度的 RF/PW 与角度正余弦空间工作。DBSCAN、HDBSCAN、Mean Shift 自动给出簇或噪声；K-means、Spectral、Hierarchical 需要的簇数统一由无监督 HDBSCAN 估计，所耗时间包含在推理过程里。聚类不接收真实源数。Spectral 的 k=1 分支会直接给一个簇。")
heading("三个深度基线", 2)
para("train.py 只为 GRU、TCN、Transformer 训练公共的六维脉冲表示学习任务。每个训练块至少有两个可见源。relation_loss 将同源脉冲对的余弦距离压低，将异源脉冲对距离推过 0.8 间隔；两类对分别平均，避免对数不平衡。scan 与 stare 各自只用 train 训练、val 选择最低验证损失轮次。局部源编号只用于同一块内的同源关系，不当作跨文件全局类别。")
para("GRU 是两层双向循环网络；TCN 是六个膨胀卷积残差块；Transformer 是四层、四注意力头的编码器。三者都输出 32 维单位向量，推理时按块加载已训练权重，再用 HDBSCAN 将向量变成最终簇标签。缺权重会返回未评估；正式运行还要求 checkpoint 标为 formal。小规模 smoke 权重仅验证代码路径。")

p = heading("五  十篇论文方法的代码路径")
p.paragraph_format.page_break_before = True
para("下表解释当前实现做了什么。它是代码导读，不表示十篇论文均已严格复现；每篇的公式对应和待核验数值见 docs/论文实现对应表.md。")
table(["注册名", "当前算法路径", "使用边界"], [
    ("edpc", "RF/PW/AOA 样本内归一化；势能过滤、自然邻域、密度峰与二阶回归预测上界找中心；再分配与并簇。", "无中心回退和阈值容差为工程补全。"),
    ("hypergraph", "去自势能过滤、密度与反向邻居选核心；模糊 C 均值结合肘部/轮廓系数选簇数；建超图并传播标签。", "选簇细节与阈值仍需原文核验。"),
    ("sapd", "以 ToA 和 PRI 为 Kalman 状态；逐候选扩展检测/漏检假设，门控、打分并做 N 扫描剪枝。", "候选或假设超上限时明确失败。"),
    ("dwell_switch", "CDIF 提 PRI 候选；双向序列搜索、FFT 判别 D&S，建立互补兼容图并取最大团组合。", "CDIF 前端是公开的实现选择。"),
    ("incremental", "按时间顺序维护 leader 中心，匹配则指数更新，否则新建；确认期后删除不足六脉冲的簇。", "代表论文的 leader 方法，特征适配有差异。"),
    ("cascaded", "自然邻域与 RLOF 过滤、适应半径 DBSCAN、邻域并簇、PRI 相位直方图找序列并去重。", "部分 RLOF 与峰阈值为显式默认值。"),
    ("grammar", "将外部子组与组语法展开成 ToA 模板，匹配启动序列、允许缺失与杂波，冲突目标拒识。", "需独立给出语法先验；当前 TSRD 无现成先验。"),
    ("dae", "二值时间格输入自编码器，输出一个已知目标的检测结果。", "需目标先验和单独训练权重。"),
    ("nmt", "DTOA 经双层双向 LSTM 逐脉冲分类；末脉冲无下一差分，标 -1。", "需跨样本一致类别语义和权重。"),
    ("pcid", "五维 PDW 与逐脉冲 128 维真实波形相关特征融合，再分类。", "TSRD 无波形相关；网络连接尚待原图核验。"),
], [2.5, 10.1, 4.0])
para("train_paper.py 为 DAE、NMT、PCID 提供独立先验 NPZ 的训练入口，要求 train/val 文件不相交且配置说明 prior_provenance。不会把 TSRD 不同文件里的局部源编号自动拼成一个全局类别。grammar 必须由调用方提供有来源说明的子组脉冲数、PRI 序列和组间间隔。")

heading("六  评分是怎样算出来的")
heading("逐样本指标", 2)
para("evaluation/metrics.py 的 evaluate 首先检查预测标签长度与输入 PDW 数相同。真实 y=-1 的虚警不参加发射源聚类主指标；被算法拒识的真实脉冲仍记错。主指标分开报告 ARI 和宏平均源匹配 F1。ARI 在真实脉冲上计算，并将每个被拒识的真实脉冲当作独立未归属项，避免所有 -1 组成虚假的大簇。真实划分退化时 ARI/AMI 留空。")
para("宏平均源 F1 先构造真实源与预测簇的交集计数，用 2×交集/(真实源大小+预测簇大小) 得到每个配对的 F1。预测簇大小包含被错误吸入的虚警，拒识真实脉冲保留在真实源大小中。匈牙利算法做一对一最大 F1 匹配；未匹配的真实源得 0，再对真实源等权平均。匹配准确率另用最大正确脉冲数的一对一匹配，不能沿用 F1 的配对。")
table(["指标组", "代码口径"], [
    ("聚类", "ARI、AMI、同质性、完整性、V-measure、FMI；来源为真实脉冲。"),
    ("脉冲对", "由源×簇交集计数得到 TP/FP/FN/TN，再算 Precision、Recall、F1、MCC，不生成 N×N 矩阵。"),
    ("源级", "源平均 F1、匹配准确率、源数绝对误差；F1≥0.5 的一对一配对计为源检出。"),
    ("错误类型", "一个真实源分到多个预测簇记过分裂；一个预测簇吸收多个真实源记错误合并。"),
    ("虚警与拒识", "真实虚警且预测 -1 为虚警 TP；真实脉冲预测 -1 为误拒；真实虚警被接收为 FN。"),
    ("父事件等权", "同 parent 的片段各获 1/片段数权重，另算匹配 F1 与准确率。"),
], [3.2, 13.4])
para("空窗口和真实脉冲少于两条时，主要聚类指标留空，并保存 validity 原因。只有一个真实源等退化真值划分时 ARI/AMI 也留空。若没有真实虚警，虚警检测分数留空。TSRD 官方兼容 F1/MCC 需显式开启 tsrd_compatibility，且用 tdc_min_best 前缀单列，不与这里的脉冲对 F1/MCC 混用。")
heading("隔离执行与失败状态", 2)
para("evaluation/runner.py 在独立进程执行算法，主进程限制单样本脉冲数、执行秒数和峰值 RSS。超出限制会记 resource_limit、timeout 或 memory_limit；缺先验是 not_evaluated；场景少于两源则可记 not_applicable。成功返回标签后才算指标并保存 prediction。不会为通过限制而暗中缩减脉冲数。")

heading("七  实验配置与结果怎么读")
para("benchmark.py 支持 inspect、train、evaluate、report、verify-sources。configs/smoke.json 用一个文件和少量窗口做本机功能验证；configs/scenarios_stare_smoke.json 走完整场景网格；configs/dense_smoke.json 检查最大源数记录；configs/full_gpu.json 是待运行的正式训练与评估设置。评估子集可由 --subsets 改成任一子集、多选或 all，但 smoke 的 files 限制会继续生效。")
para("每个运行目录中的 config_snapshot.json 保存配置、代码、软件版本、profile 与权重校验；samples 下的同一 sample_id 对应一份派生 PDW/真值/parent，predictions 保存预测标签，rows 保存每个算法的状态、参数、原因和指标。report.py 把 rows 展平成 per_sample.csv，再按子集、算法、场景、track、fidelity 汇总到 summary.csv/json。汇总先在原始文件内部平均，再让文件等权；多种子给标准差，至少两个有效文件才给按文件 bootstrap 的置信区间。")
para("读成绩时先查看 status 与 metric_validity_counts，再看 ARI 和源平均 F1。成功执行但窗口为空，不能把成功数当成有效分数。当前保存的 smoke 结果是功能验证，尚不能作为全部 TSRD、三个种子或十篇论文严格复现后的排名。源码和来源核对文件位于 docs/本机验证记录.md、docs/论文实现对应表.md；原论文与 TSRD 的保护核验可由 verify-sources 重跑。")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "雷达分选 Benchmark 代码逻辑导读"
doc.core_properties.subject = "数据处理 算法 实验 评估"
doc.core_properties.author = ""
doc.save(OUTPUT)
print(OUTPUT)
